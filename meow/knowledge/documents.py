"""Files you keep - Phase 2.4.

Everything Meow has made until now died with the session: a sentence spoken, a
button pressed, a note typed into whatever happened to have focus. A spreadsheet
on the desktop is the first thing that outlives it.

Three formats, one shape each. Deliberately plain: a Word document with a
heading and paragraphs, a spreadsheet with a header row and data, a deck with a
title and bullets. Anything more elaborate is a different request, and the model
filling these in has to be able to describe the whole thing in one tool call.

**Nothing is ever silently overwritten.** A name that already exists gets a
number, so "notes.docx" becomes "notes 2.docx". Losing work to an assistant
being helpful is the worst thing in this file, and the check costs nothing.

**Everything lands in one folder.** `Documents/Meow` rather than wherever the
working directory happens to be, because a file the user cannot find is barely
better than no file. The path comes back with the result so the cat can say it
out loud.
"""

from __future__ import annotations

import os
import re
from dataclasses import dataclass
from pathlib import Path

# Windows forbids these in a file name, and a model asked for a title will
# cheerfully produce "Q3: profit/loss".
_ILLEGAL = re.compile(r'[<>:"/\\|?*\x00-\x1f]')

MAX_NAME_LENGTH = 80


@dataclass(frozen=True)
class Document:
    path: Path
    kind: str
    detail: str

    def describe(self) -> str:
        return f"{self.detail}, saved as {self.path.name}"


def output_folder() -> Path:
    """Where files go. Made if it does not exist.

    One predictable place beats the working directory, which for a voice
    assistant is wherever it happened to be launched from.
    """
    folder = Path(os.environ.get("USERPROFILE", Path.home())) / "Documents" / "Meow"
    folder.mkdir(parents=True, exist_ok=True)
    return folder


def safe_name(name: str, extension: str) -> str:
    """A file name Windows will accept, with the right extension."""
    cleaned = _ILLEGAL.sub("", name).strip().strip(".")
    cleaned = " ".join(cleaned.split()) or "untitled"
    cleaned = cleaned[:MAX_NAME_LENGTH]
    if not cleaned.lower().endswith(extension):
        cleaned += extension
    return cleaned


def unused_path(name: str, extension: str) -> Path:
    """A path that does not exist yet.

    Numbers rather than overwrites. An assistant that quietly replaces a file
    the user spent an hour on has done more damage than one that refuses to
    write at all.
    """
    folder = output_folder()
    candidate = folder / safe_name(name, extension)
    if not candidate.exists():
        return candidate

    stem = candidate.stem
    for index in range(2, 100):
        numbered = folder / f"{stem} {index}{extension}"
        if not numbered.exists():
            return numbered
    # Ninety-eight files of the same name is not a case worth handling well.
    return folder / f"{stem} {os.getpid()}{extension}"


def make_docx(name: str, heading: str, paragraphs: list[str]) -> Document:
    """A Word document: one heading, then paragraphs."""
    from docx import Document as WordDocument

    document = WordDocument()
    if heading.strip():
        document.add_heading(heading.strip(), level=1)
    for text in paragraphs:
        if str(text).strip():
            document.add_paragraph(str(text).strip())

    path = unused_path(name or heading or "document", ".docx")
    document.save(str(path))
    return Document(path, "document",
                    f"{len(paragraphs)} paragraphs about {heading or name}")


def make_xlsx(name: str, headers: list[str], rows: list[list[str]],
              sheet_title: str = "Sheet1") -> Document:
    """A spreadsheet: a bold header row, then data, columns sized to fit."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = safe_name(sheet_title, "")[:31] or "Sheet1"

    if headers:
        sheet.append([str(header) for header in headers])
        for cell in sheet[1]:
            cell.font = Font(bold=True)
    for row in rows:
        sheet.append([str(value) for value in row])

    # Width from the longest value in each column. A spreadsheet where every
    # column shows ##### is not usable, and nobody wants to be told to widen
    # the columns by the thing that made them.
    for index, column in enumerate(sheet.columns, start=1):
        longest = max((len(str(cell.value or "")) for cell in column), default=8)
        sheet.column_dimensions[
            sheet.cell(row=1, column=index).column_letter
        ].width = min(60, max(10, longest + 2))

    path = unused_path(name or sheet_title or "spreadsheet", ".xlsx")
    workbook.save(str(path))
    return Document(path, "spreadsheet",
                    f"{len(rows)} rows across {len(headers)} columns")


def make_pptx(name: str, title: str, slides: list[dict]) -> Document:
    """A deck: a title slide, then one slide per entry.

    Each entry is {"title": ..., "bullets": [...]}. Layout 1 is the built-in
    title-and-content layout, which every template has - picking a fancier one
    breaks the moment someone applies a theme.
    """
    from pptx import Presentation

    presentation = Presentation()

    opening = presentation.slides.add_slide(presentation.slide_layouts[0])
    opening.shapes.title.text = title or name or "Untitled"
    if len(opening.placeholders) > 1:
        opening.placeholders[1].text = "made by meow"

    for entry in slides:
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = str(entry.get("title", "")).strip()
        bullets = entry.get("bullets") or []
        frame = slide.placeholders[1].text_frame
        frame.clear()
        for position, bullet in enumerate(bullets):
            paragraph = frame.paragraphs[0] if position == 0 else frame.add_paragraph()
            paragraph.text = str(bullet).strip()

    path = unused_path(name or title or "slides", ".pptx")
    presentation.save(str(path))
    return Document(path, "slides", f"{len(slides)} slides about {title or name}")


def open_document(document: Document) -> bool:
    """Open it in whatever Windows uses for that type."""
    try:
        os.startfile(str(document.path))  # noqa: S606 - the point of this
        return True
    except OSError:
        return False
